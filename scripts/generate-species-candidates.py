"""Generate beta.3 candidate Species Items from the official 2026 CRD DOCX.

This developer-only extractor writes the private review source only. It never
creates a system Compendium or changes system.json.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import unicodedata
from pathlib import Path

from docx import Document


SCHEMA_VERSION = 1
BOOK = "Cypher Reference Document (2026)"
LICENSE = "2026 Cypher Open License"
SPECIES_ICON = "systems/cypherv2/assets/icons/cypherspecies.png"
SKILL_ICON = "systems/cypherv2/assets/icons/cypherskill.png"
SKILL_PREFIX = "Compendium.cypherv2.skills.Item."

FANTASY_NAMES = [
    "Dragonfolk", "Dwarf", "Elf", "Gnome", "Halfling", "Hellborn", "Human", "Orc",
]
SCIENCE_FICTION_NAMES = [
    "Aarak", "Cyborg", "Delph", "D’nec", "Drakain", "Human", "Mutant", "Naron",
    "Prota", "Rigellian", "Stelan", "Vendeer", "Zantari",
]
SPECIES_NAMES = FANTASY_NAMES + [name for name in SCIENCE_FICTION_NAMES if name != "Human"]

FIXED_SKILLS = {
    "Dragonfolk": [("Intimidation", "<p>Except against other dragonfolk and dragons.</p>")],
    "Dwarf": [("Navigation", "<p>Only underground or in mountainous areas.</p>")],
    "Elf": [("Stealth", "<p>Only in forested areas.</p>")],
    "Gnome": [("Crafting", "")],
    "Halfling": [("Escaping", "")],
    "Hellborn": [("Magic Lore", "")],
    "Orc": [("Intimidation", "<p>Except against other orcs.</p>")],
    "Cyborg": [("Hacking", "")],
    "D’nec": [("Systems Operation", "")],
    "Drakain": [("Intimidation", "<p>Except against other drakain.</p>")],
    "Naron": [("Recognizing Motive", "")],
    "Prota": [("Charm", "")],
    "Rigellian": [("Systems Operation", "")],
    "Stelan": [("Gathering information", "")],
    "Vendeer": [("Perception", "")],
    "Zantari": [("Perception", "")],
}

WOUND_BONUSES = {
    "Orc": {"minor": 1, "moderate": 0, "major": 0},
    "Drakain": {"minor": 1, "moderate": 0, "major": 0},
    "Prota": {"minor": 1, "moderate": 0, "major": 0},
}


def normalized_space(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def slugify(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def stable_id(kind: str, key: str) -> str:
    return hashlib.sha256(f"cypherv2-beta3:{kind}:{key}".encode("utf-8")).hexdigest()[:16]


def provenance() -> dict:
    return {
        "kind": "other", "sourceUuid": "", "instanceId": "", "grantId": "",
        "status": "active", "contentUuid": "", "contentKey": "",
        "replacement": {
            "active": False, "originalName": "", "originalContentUuid": "",
            "originalContentKey": "", "replacementName": "",
            "replacementContentUuid": "", "replacementContentKey": "",
            "selectionKind": "none",
        },
    }


def package_instance() -> dict:
    return {
        "sourceUuid": "", "instanceId": "", "role": "primary", "attachedAt": 0,
        "selections": {
            "edgePool": "none", "superheroicsPool": "none", "powerShifts": [],
            "poolChoices": [], "skillChoices": [], "abilityChoices": [],
            "descriptorChoices": [], "suppressedGrantIds": [],
        },
        "parent": provenance(),
    }


def paragraph_page_numbers(paragraphs) -> list[int]:
    page = 1
    result = []
    for paragraph in paragraphs:
        properties = paragraph._p.pPr
        if properties is not None and properties.pageBreakBefore is not None:
            page += 1
        result.append(page)
        page += len(paragraph._p.xpath(".//w:lastRenderedPageBreak"))
        page += len(paragraph._p.xpath('.//w:br[@w:type="page"]'))
    return result


def runs_html(paragraph) -> str:
    fragments = []
    for run in paragraph.runs:
        if not run.text:
            continue
        value = html.escape(run.text.replace("\u2028", " ").replace("\u2029", " ")).replace("\n", "<br>")
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        fragments.append(value)
    return "".join(fragments).strip()


def section_html(paragraphs) -> str:
    fragments = []
    list_items = []

    def flush_list() -> None:
        if list_items:
            fragments.append(f"<ul>{''.join(list_items)}</ul>")
            list_items.clear()

    for paragraph in paragraphs:
        body = runs_html(paragraph)
        if not body:
            continue
        if paragraph.style.name == "Compact":
            list_items.append(f"<li><p>{body}</p></li>")
            continue
        flush_list()
        if paragraph.style.name == "Callout":
            fragments.append(f'<aside class="species-callout"><p>{body}</p></aside>')
        else:
            fragments.append(f"<p>{body}</p>")
    flush_list()
    return "".join(fragments)


def extract_species_sections(document: Document) -> tuple[dict[str, dict], list[int], dict[str, list[str]]]:
    paragraphs = document.paragraphs
    pages = paragraph_page_numbers(paragraphs)

    def heading_index(title: str, style: str) -> int:
        matches = [i for i, paragraph in enumerate(paragraphs)
                   if normalized_space(paragraph.text) == title and paragraph.style.name == style]
        if len(matches) != 1:
            raise ValueError(f"Expected one {style} heading '{title}', found {len(matches)}.")
        return matches[0]

    fantasy_start = heading_index("Fantasy Species", "Heading 3")
    fantasy_end = next(i for i in range(fantasy_start + 1, len(paragraphs)) if paragraphs[i].style.name == "H2")
    science_start = heading_index("Science Fiction Species", "Heading 3")
    science_end = heading_index("Species or Type?", "Heading 3")

    def collect(start: int, end: int, style: str) -> list[dict]:
        headings = [i for i in range(start + 1, end) if paragraphs[i].style.name == style]
        result = []
        for position, index in enumerate(headings):
            name = normalized_space(paragraphs[index].text)
            next_index = headings[position + 1] if position + 1 < len(headings) else end
            result.append({"name": name, "heading": index, "end": next_index})
        return result

    fantasy = collect(fantasy_start, fantasy_end, "Type Headers")
    science = collect(science_start, science_end, "Heading 3")
    fantasy_names = [entry["name"] for entry in fantasy]
    science_names = [entry["name"] for entry in science]
    if fantasy_names != FANTASY_NAMES:
        raise ValueError(f"Fantasy Species inventory mismatch: {fantasy_names}")
    if science_names != SCIENCE_FICTION_NAMES:
        raise ValueError(f"Science Fiction Species inventory mismatch: {science_names}")

    sections = {entry["name"]: entry for entry in fantasy}
    for entry in science:
        sections.setdefault(entry["name"], entry)
    return sections, pages, {"fantasy": fantasy_names, "scienceFiction": science_names}


def load_public_references(reference_path: Path) -> dict[str, dict]:
    source = json.loads(reference_path.read_text(encoding="utf-8"))
    skills = {
        document["name"]: document
        for document in source["packs"]["skills"]["documents"]
        if document.get("type") == "skill"
    }
    required_skills = {name for grants in FIXED_SKILLS.values() for name, _ in grants} | {"Charm", "Deception"}
    missing = sorted(required_skills - skills.keys())
    if missing:
        raise ValueError(f"Public Skills source is missing: {', '.join(missing)}")
    return skills


def snapshot(document: dict, fallback_img: str) -> dict:
    return {
        "name": document["name"],
        "img": document.get("img") or fallback_img,
        "system": document["system"],
    }


def skill_grants(species_name: str, skills: dict[str, dict]) -> list[dict]:
    grants = []
    for skill_name, notes in FIXED_SKILLS.get(species_name, []):
        source = skills[skill_name]
        grants.append({
            "id": stable_id("species-skill-grant", f"{slugify(species_name)}:{source['_id']}"),
            "skillUuid": f"{SKILL_PREFIX}{source['_id']}",
            "customName": "", "notes": notes, "rank": "trained",
            "snapshot": snapshot(source, SKILL_ICON), "alternatives": [],
        })
    return grants


def skill_choice_groups(species_name: str, skills: dict[str, dict]) -> list[dict]:
    if species_name != "Naron":
        return []
    options = []
    for skill_name in ("Charm", "Deception"):
        source = skills[skill_name]
        options.append({
            "id": stable_id("species-skill-option", f"naron:{source['_id']}"),
            "skillUuid": f"{SKILL_PREFIX}{source['_id']}",
            "customName": "", "notes": "", "snapshot": snapshot(source, SKILL_ICON),
        })
    return [{
        "id": stable_id("species-skill-choice", "naron:charm-or-deception"),
        "choose": 1, "rank": "trained", "options": options,
    }]


def descriptor_choice_groups(species_name: str) -> list[dict]:
    if species_name != "Human":
        return []
    return [{
        "id": stable_id("species-descriptor-choice", "human:second-descriptor"),
        "choose": 1,
        "sourceMode": "catalog",
        "catalogItemType": "descriptor",
        "options": [],
    }]


def build_species_documents(document: Document, references_path: Path) -> tuple[list[dict], dict[str, list[str]]]:
    paragraphs = document.paragraphs
    sections, pages, inventory = extract_species_sections(document)
    skills = load_public_references(references_path)
    documents = []
    for sort, name in enumerate(SPECIES_NAMES):
        section = sections[name]
        source_paragraphs = paragraphs[section["heading"] + 1:section["end"]]
        slug = slugify(name)
        item_id = stable_id("species", slug)
        genre_tags = ["fantasy", "science-fiction"] if name == "Human" else (["fantasy"] if name in FANTASY_NAMES else ["science-fiction"])
        source_page = str(pages[section["heading"]])
        if name == "Human":
            source_page = "103; 158–159"
        system = {
            "schemaVersion": 1,
            "slug": slug,
            "description": section_html(source_paragraphs),
            "source": {"uuid": "", "book": BOOK, "page": source_page, "license": LICENSE},
            "automation": {
                "mode": "descriptive",
                "duration": {"enabled": False, "trigger": "recovery"},
                "rollDefaults": {},
            },
            "ruleElements": [],
            "tags": ["species", *genre_tags],
            "grantedBy": provenance(),
            "poolBonuses": {"might": 0, "speed": 0, "intellect": 0},
            "woundBonuses": WOUND_BONUSES.get(name, {"minor": 0, "moderate": 0, "major": 0}),
            "edgeGrant": {"mode": "none", "pool": "none", "amount": 1},
            "weaponUse": {"light": False, "medium": False, "heavy": False},
            "weaponFamilies": [],
            "armorUse": {"light": False, "medium": False, "heavy": False},
            "cypherLimitBonus": 0,
            "skillGrants": skill_grants(name, skills),
            "choiceGroups": skill_choice_groups(name, skills),
            "abilityGrants": [],
            "abilityChoiceGroups": [],
            "descriptorGrants": [],
            "descriptorChoiceGroups": descriptor_choice_groups(name),
            "instance": package_instance(),
        }
        documents.append({
            "name": name, "type": "species", "_id": item_id, "img": SPECIES_ICON,
            "system": system, "effects": [], "folder": None, "sort": sort * 1000,
            "ownership": {"default": 0}, "flags": {},
            "_stats": {
                "coreVersion": "14.360", "systemId": None, "systemVersion": None,
                "createdTime": None, "modifiedTime": None, "lastModifiedBy": None,
                "compendiumSource": None, "duplicateSource": None, "exportSource": None,
            },
            "_key": f"!items!{item_id}",
        })
    return documents, inventory


def build_candidate(document_path: Path, references_path: Path) -> dict:
    document = Document(document_path)
    species, inventory = build_species_documents(document, references_path)
    return {
        "$schema": "./schema.json",
        "schemaVersion": SCHEMA_VERSION,
        "source": {
            "kind": "official-reference-document-extraction",
            "title": BOOK,
            "file": document_path.name,
            "sha256": hashlib.sha256(document_path.read_bytes()).hexdigest(),
            "license": LICENSE,
            "note": "Candidate content only; requires in-Foundry review before promotion.",
        },
        "target": {"systemId": "cypherv2", "futurePackId": "species"},
        "inventory": inventory,
        "counts": {"species": len(species)},
        "review": {
            "status": "candidate",
            "manualTreatment": [
                "Dwarf recovery bonus remains descriptive; no RecoveryService automation is introduced.",
                "D’nec additional Intellect recovery remains descriptive; no targeted recovery automation is introduced.",
                "Mutant's adapted tier 1 Focus Ability remains a manual player/GM choice with no Ability catalog.",
                "Natural-1 exceptions, conditional damage, resistance, movement, telepathy, Assets, and turn denial remain descriptive.",
            ],
            "human": "One shared Human document covers both Fantasy and Science Fiction appearances and offers one dynamic all-available-Descriptors choice.",
        },
        "species": species,
    }


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=project_root / "_reference" / "Cypher-Reference-Document-2026-07-29.docx")
    parser.add_argument("--references", type=Path, default=project_root / "content" / "core-items" / "reviewed-packs.json")
    parser.add_argument("--output", type=Path, default=project_root / "content" / "species" / "candidate-pack.json")
    args = parser.parse_args()
    candidate = build_candidate(args.source, args.references)
    if candidate["counts"] != {"species": 20}:
        raise ValueError(f"Unexpected candidate counts: {candidate['counts']}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(candidate, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), **candidate["counts"]}, indent=2))


if __name__ == "__main__":
    main()

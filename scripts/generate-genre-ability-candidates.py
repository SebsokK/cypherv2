"""Generate beta.3 candidate Genre Ability Items from the official 2026 CRD.

The generated JSON is a private editorial source.  It is deliberately separate
from both the private Foundry Working pack and any future public Compendium.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import unicodedata
from pathlib import Path

from docx import Document


SCHEMA_VERSION = 1
BOOK = "Cypher Reference Document (2026)"
LICENSE = "2026 Cypher Open License"
ABILITY_ICON = "systems/cypherv2/assets/icons/cypherability.png"
ABILITY_PREFIX = "Compendium.cypherv2.genre-abilities.Item."

REAL_WORLD_MID = ["Additional Skill (Tier 3)"]
REAL_WORLD_HIGH = ["Additional Skill (Tier 6)"]
REAL_WORLD_IDS = {
    "Additional Skill (Tier 3)": "fREdI45SXu0WPeeS",
    "Additional Skill (Tier 6)": "xXSq03yZHRS978cj",
}

FANTASY_MID = [
    "A Bit of Magic", "Cypher Use", "Danger Instinct", "Disappear Into Shadow",
    "Discerning Mind", "Elemental Protection", "Enhanced Stat", "Exceptional Follower",
    "From the Shadows", "Fury", "Pry Open Defense", "Puncturing Attack", "Snipe",
    "Strategize", "Tough", "Winning Smile",
]
FANTASY_HIGH = [
    "Assassin Strike", "Concussive Force", "Inspire Action", "Invisibility", "Jump Attack",
    "Magic Portal", "Mask", "Resilience", "Spellbreaker", "Spin Attack", "Will of a Leader",
]
SCIENCE_FICTION_MID = [
    "Black Thumb", "Cypher Use", "Disable Mechanism", "Enhanced Stat", "Exceptional Follower",
    "Hands on the Wheel", "Incredible Health", "Machine Companion", "Mind Reading", "Snipe", "Spray",
]
SCIENCE_FICTION_HIGH = [
    "Arc Spray", "Improved Machine Companion", "Inspire Action", "Knowledge Expert",
    "Lethal Capability", "Severe Machine Disruption", "Technology Expert", "Telepathic Network",
]
ORIGIN = [
    "Adhesive Mobility", "Amazing Invulnerability", "Amazing Tools", "Armored Body",
    "Astonishing Teleport", "Awesome Force Field", "Duplicate", "Extraordinary Leap",
    "Fantastic Armament", "Fantastic Vehicle", "Incredible Instinct", "Incredible Velocity",
    "Intangible", "Invisible Knack", "Power Cypher Use", "Powerful Blast",
    "Regenerative Healing", "Shrink", "Skill Exemplar", "Stretchy", "Superhero Versatility",
    "Team-Up Ally", "Telepathic Prodigy", "Unbelievable Transformation", "Uncanny Flight",
    "Unyielding Shield",
]

PROGRESSION_ORDER = FANTASY_MID + FANTASY_HIGH + [
    name for name in SCIENCE_FICTION_MID + SCIENCE_FICTION_HIGH
    if name not in FANTASY_MID and name not in FANTASY_HIGH
]

SECTION_CONFIGS = [
    ("Mid-Tier Fantasy Abilities", "fantasy", "mid-tier", FANTASY_MID),
    ("High-Tier Fantasy Abilities", "fantasy", "high-tier", FANTASY_HIGH),
    ("Mid-Tier Science Fiction Abilities", "scienceFiction", "mid-tier", SCIENCE_FICTION_MID),
    ("High-Tier Science Fiction Abilities", "scienceFiction", "high-tier", SCIENCE_FICTION_HIGH),
    ("Origin Superhero Abilities", "superhero", "origin", ORIGIN),
]

ACTIVATIONS = {
    "Additional Skill (Tier 3)": "passive", "Additional Skill (Tier 6)": "passive",
    "A Bit of Magic": "special", "Cypher Use": "passive", "Danger Instinct": "enabler",
    "Disappear Into Shadow": "enabler", "Discerning Mind": "enabler",
    "Elemental Protection": "action", "Enhanced Stat": "passive",
    "Exceptional Follower": "passive", "From the Shadows": "enabler", "Fury": "action",
    "Pry Open Defense": "action", "Puncturing Attack": "passive", "Snipe": "firstAction",
    "Strategize": "timed", "Tough": "passive", "Winning Smile": "action",
    "Assassin Strike": "enabler", "Concussive Force": "action", "Inspire Action": "action",
    "Invisibility": "action", "Jump Attack": "action", "Magic Portal": "action",
    "Mask": "action", "Resilience": "passive", "Spellbreaker": "special",
    "Spin Attack": "action", "Will of a Leader": "timed", "Black Thumb": "passive",
    "Disable Mechanism": "action", "Hands on the Wheel": "enabler",
    "Incredible Health": "passive", "Machine Companion": "passive", "Mind Reading": "action",
    "Spray": "action", "Arc Spray": "action", "Improved Machine Companion": "passive",
    "Knowledge Expert": "enabler", "Lethal Capability": "passive",
    "Severe Machine Disruption": "action", "Technology Expert": "passive",
    "Telepathic Network": "special", "Adhesive Mobility": "passive",
    "Amazing Invulnerability": "passive", "Amazing Tools": "action", "Armored Body": "passive",
    "Astonishing Teleport": "special", "Awesome Force Field": "action", "Duplicate": "action",
    "Extraordinary Leap": "passive", "Fantastic Armament": "special", "Fantastic Vehicle": "special",
    "Incredible Instinct": "passive", "Incredible Velocity": "special", "Intangible": "special",
    "Invisible Knack": "special", "Power Cypher Use": "passive", "Powerful Blast": "special",
    "Regenerative Healing": "perpetual", "Shrink": "firstAction", "Skill Exemplar": "passive",
    "Stretchy": "special", "Superhero Versatility": "passive", "Team-Up Ally": "passive",
    "Telepathic Prodigy": "enabler", "Unbelievable Transformation": "lastAction",
    "Uncanny Flight": "passive", "Unyielding Shield": "passive",
}

# Only unconditional activation costs are represented. Conditional alternative
# costs remain in the authoritative rich text rather than being flattened.
EXPLICIT_COSTS = {
    "Danger Instinct": (3, False, ["speed"]), "Discerning Mind": (2, False, ["intellect"]),
    "Elemental Protection": (4, True, ["intellect"]), "Fury": (3, False, ["might"]),
    "Pry Open Defense": (4, False, ["intellect"]), "Snipe": (2, False, ["speed"]),
    "Strategize": (4, False, ["intellect"]), "Winning Smile": (2, True, ["intellect"]),
    "Assassin Strike": (5, False, ["speed"]), "Concussive Force": (7, False, ["intellect"]),
    "Inspire Action": (4, False, ["intellect"]), "Invisibility": (4, True, ["intellect"]),
    "Jump Attack": (5, True, ["might"]), "Magic Portal": (6, False, ["intellect"]),
    "Mask": (5, False, ["intellect"]), "Spellbreaker": (4, True, ["intellect"]),
    "Spin Attack": (5, True, ["speed"]), "Will of a Leader": (9, False, ["intellect"]),
    "Disable Mechanism": (3, True, ["intellect"]), "Hands on the Wheel": (2, True, ["intellect"]),
    "Mind Reading": (2, False, ["intellect"]), "Spray": (2, False, ["speed"]),
    "Arc Spray": (3, False, ["speed"]), "Severe Machine Disruption": (5, True, ["intellect"]),
    "Telepathic Network": (0, True, ["intellect"]), "Amazing Tools": (2, False, ["intellect"]),
    "Duplicate": (2, True, ["might"]), "Shrink": (1, True, ["might"]),
}

CONDITIONAL_COSTS = {
    "A Bit of Magic", "Awesome Force Field", "Fantastic Armament",
    "Fantastic Vehicle", "Regenerative Healing",
}

ROLL_METADATA = {
    "Pry Open Defense": ("attack", 0, "long", "single"),
    "Concussive Force": ("attack", 5, "long", "multiple"),
    "Spellbreaker": ("attack", 0, "immediate", "single"),
    "Disable Mechanism": ("attack", 0, "immediate", "single"),
    "Severe Machine Disruption": ("attack", 0, "immediate", "single"),
    "Powerful Blast": ("attack", 4, "long", "single"),
}

DURATIONS = {
    "Elemental Protection": "1-hour-or-longer", "Fury": "10-minute-or-longer",
    "Pry Open Defense": "10-minute-or-longer", "Strategize": "10-hour",
    "Winning Smile": "10-minute-or-longer", "Invisibility": "10-minute-or-longer",
    "Magic Portal": "10-minute-or-longer", "Spellbreaker": "10-minute-or-longer",
    "Will of a Leader": "1-hour-or-longer", "Hands on the Wheel": "10-minute-or-longer",
    "Mind Reading": "10-minute-or-longer", "Awesome Force Field": "10-minute-or-longer",
    "Duplicate": "10-minute-or-longer", "Shrink": "10-minute-or-longer",
}

COLLISIONS = [
    ("Cypher Use", "B", ["type"], "Genre wording includes repeat acquisition and replacement context."),
    ("Danger Instinct", "A", ["focus"], "Same surprise-movement and eased-defense mechanics."),
    ("Tough", "A", ["focus"], "Same additional minor-wound capacity."),
    ("Invisibility", "B", ["focus"], "Genre duration and target-ending wording belongs to its own acquisition context."),
    ("Jump Attack", "A", ["focus"], "Same jump, attack, knockdown, and Effort mechanics."),
    ("Resilience", "A", ["focus"], "Same additional moderate-wound capacity."),
    ("Disable Mechanism", "A", ["focus"], "Same machine-disruption choices and Effort extensions."),
    ("Machine Companion", "B", ["focus"], "Genre text includes the existing-companion upgrade alternative."),
    ("Mind Reading", "A", ["type", "focus"], "Type occurrence matches; Genre keeps independent provenance and catalog identity."),
    ("Spray", "A", ["focus"], "Same rapid-fire asset, ammunition, and damage tradeoff."),
    ("Arc Spray", "A", ["focus"], "Same three-target hindered rapid-fire attack."),
    ("Improved Machine Companion", "A", ["focus"], "Same level 5 companion and upgrade alternative."),
    ("Lethal Capability", "A", ["focus"], "At least one Focus occurrence matches; other Focus uses remain acquisition-specific."),
    ("Telepathic Network", "B", ["focus"], "Genre wording and scalable network expansion remain independently sourced."),
]


def normalized_space(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def slugify(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def stable_id(kind: str, key: str) -> str:
    return hashlib.sha256(f"cypherv2-beta3:{kind}:{key}".encode()).hexdigest()[:16]


def paragraph_page_numbers(paragraphs) -> list[int]:
    page = 1
    result = []
    for paragraph in paragraphs:
        properties = paragraph._p.pPr
        if properties is not None and properties.pageBreakBefore is not None:
            page += 1
        result.append(page)
        page += len(paragraph._p.xpath(".//w:lastRenderedPageBreak"))
        page += len(paragraph._p.xpath('.//w:br[@w:type="page"]'))
    return result


def leading_bold(paragraph) -> str:
    parts = []
    for run in paragraph.runs:
        if run.bold:
            parts.append(run.text)
        elif parts:
            break
    return "".join(parts).strip()


def runs_html(paragraph, remove_prefix: str = "") -> str:
    remaining = len(remove_prefix)
    fragments = []
    for run in paragraph.runs:
        text = run.text
        if remaining:
            consumed = min(remaining, len(text))
            text = text[consumed:]
            remaining -= consumed
        if not text:
            continue
        value = html.escape(text.replace("\u2028", " ").replace("\u2029", " ")).replace("\n", "<br>")
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        fragments.append(value)
    return "".join(fragments).strip()


def paragraphs_html(paragraphs, first_prefix: str) -> str:
    fragments = []
    for index, paragraph in enumerate(paragraphs):
        body = runs_html(paragraph, first_prefix if index == 0 else "")
        if not body:
            continue
        if paragraph.style.name == "Compact":
            fragments.append(f"<p class=\"compact\">{body}</p>")
        elif paragraph.style.name == "Callout":
            fragments.append(f"<aside><p>{body}</p></aside>")
        else:
            fragments.append(f"<p>{body}</p>")
    return "".join(fragments)


def provenance() -> dict:
    return {
        "kind": "other", "sourceUuid": "", "instanceId": "", "grantId": "",
        "status": "active", "contentUuid": "", "contentKey": "",
        "replacement": {
            "active": False, "originalName": "", "originalContentUuid": "", "originalContentKey": "",
            "replacementName": "", "replacementContentUuid": "", "replacementContentKey": "",
            "selectionKind": "none",
        },
    }


def extract_sections(document: Document) -> tuple[dict[str, list[dict]], list[int]]:
    paragraphs = document.paragraphs
    pages = paragraph_page_numbers(paragraphs)
    occurrences: dict[str, list[dict]] = {}
    for heading, genre, band, expected_names in SECTION_CONFIGS:
        heading_index = next((i for i, p in enumerate(paragraphs) if normalized_space(p.text) == heading), None)
        if heading_index is None:
            raise ValueError(f"Missing CRD heading: {heading}")
        end = heading_index + 1
        while end < len(paragraphs) and paragraphs[end].style.name not in ("Heading 2", "H2"):
            if band != "origin" and paragraphs[end].style.name == "Heading 3":
                break
            end += 1
        starts = []
        expected_set = set(expected_names)
        for index in range(heading_index + 1, end):
            prefix = leading_bold(paragraphs[index])
            if not prefix.endswith(":"):
                continue
            heading_text = normalized_space(prefix[:-1])
            candidate = next((name for name in expected_names if heading_text == name or heading_text.startswith(f"{name} (")), None)
            if candidate in expected_set:
                starts.append((index, candidate, prefix))
        found = [name for _, name, _ in starts]
        if found != expected_names:
            raise ValueError(f"{heading} inventory mismatch. Expected {expected_names}; found {found}.")
        for position, (start, name, prefix) in enumerate(starts):
            stop = starts[position + 1][0] if position + 1 < len(starts) else end
            occurrences.setdefault(name, []).append({
                "name": name, "genre": genre, "band": band, "page": pages[start],
                "prefix": prefix, "paragraphs": list(paragraphs[start:stop]),
            })
    return occurrences, pages


def extract_real_world_progression(document: Document, pages: list[int]) -> dict[str, list[dict]]:
    rules = [
        ("Additional Skill (Tier 3)", "mid-tier", "Genre Abilities: At tier 3,", "Genre Abilities: "),
        ("Additional Skill (Tier 6)", "high-tier", "At tier 6:", "At tier 6: "),
    ]
    result = {}
    for name, band, start, remove_prefix in rules:
        index = next(
            (i for i, paragraph in enumerate(document.paragraphs) if normalized_space(paragraph.text).startswith(start)),
            None,
        )
        if index is None:
            raise ValueError(f"Missing CRD Real World Genre Ability rule: {name}")
        result[name] = [{
            "name": name,
            "genre": "realWorld",
            "band": band,
            "page": pages[index],
            "prefix": remove_prefix,
            "paragraphs": [document.paragraphs[index]],
        }]
    return result


def base_system(slug: str, description: str, page: str, tags: list[str]) -> dict:
    return {
        "schemaVersion": 1, "slug": slug, "description": description,
        "source": {"uuid": "", "book": BOOK, "page": page, "license": LICENSE},
        "automation": {
            "mode": "descriptive", "duration": {"enabled": False, "trigger": "recovery"},
            "rollDefaults": {},
        },
        "ruleElements": [], "tags": tags, "grantedBy": provenance(),
    }


def foundry_item(name: str, item_id: str, system: dict, sort: int, metadata: dict) -> dict:
    return {
        "name": name, "type": "ability", "_id": item_id, "img": ABILITY_ICON,
        "system": system, "effects": [], "folder": None, "sort": sort,
        "ownership": {"default": 0}, "flags": {"cypherv2": {"genreAbility": metadata}},
        "_stats": {
            "coreVersion": "14.360", "systemId": None, "systemVersion": None,
            "createdTime": None, "modifiedTime": None, "lastModifiedBy": None,
            "compendiumSource": None, "duplicateSource": None, "exportSource": None,
        },
        "_key": f"!items!{item_id}",
    }


def progression_memberships(name: str) -> list[dict]:
    result = []
    for genre, mid, high in (
        ("realWorld", REAL_WORLD_MID, REAL_WORLD_HIGH),
        ("fantasy", FANTASY_MID, FANTASY_HIGH),
        ("scienceFiction", SCIENCE_FICTION_MID, SCIENCE_FICTION_HIGH),
    ):
        if name in mid:
            result.append({"genre": genre, "band": "mid-tier", "minimumTier": 3})
        if name in high:
            result.append({"genre": genre, "band": "high-tier", "minimumTier": 6})
    return result


def snapshot(document: dict) -> dict:
    return {"name": document["name"], "img": document["img"], "system": document["system"]}


def build_candidate(document_path: Path) -> dict:
    document = Document(document_path)
    occurrences, pages = extract_sections(document)
    occurrences.update(extract_real_world_progression(document, pages))
    expected = PROGRESSION_ORDER + ORIGIN + REAL_WORLD_MID + REAL_WORLD_HIGH
    if list(occurrences) != expected:
        raise ValueError("Combined Genre Ability inventory/order differs from the reviewed 69-item inventory.")

    abilities = []
    by_name = {}
    for sort, name in enumerate(expected):
        variants = occurrences[name]
        first = variants[0]
        catalog = "origin" if name in ORIGIN else "progression"
        memberships = [] if catalog == "origin" else progression_memberships(name)
        genres = ["superhero"] if catalog == "origin" else [entry["genre"] for entry in memberships]
        band = "origin" if catalog == "origin" else memberships[0]["band"]
        minimum_tier = 1 if catalog == "origin" else memberships[0]["minimumTier"]
        minimum_rank = 2 if name == "Armored Body" else 0
        pages = list(dict.fromkeys(str(variant["page"]) for variant in variants))
        slug = slugify(name)
        item_id = REAL_WORLD_IDS.get(name, stable_id("genre-ability", slug))
        description = paragraphs_html(first["paragraphs"], first["prefix"])
        amount, scalable, allowed_pools = EXPLICIT_COSTS.get(name, (0, False, []))
        roll, damage, ability_range, target_mode = ROLL_METADATA.get(name, ("none", 0, "", "none"))
        tags = ["genre-ability", f"catalog:{catalog}"]
        tags.extend(f"genre:{genre}" for genre in genres)
        if catalog == "progression":
            tags.append(f"progression:{band}")
        if name in CONDITIONAL_COSTS:
            tags.append("manual-conditional-cost")
        if name in ("Incredible Instinct", "Skill Exemplar"):
            tags.append("tier-3-improvement")
        if minimum_rank:
            tags.append(f"minimum-superhero-rank:{minimum_rank}")
        system = base_system(slug, description, ", ".join(pages), tags)
        if name in DURATIONS:
            system["automation"]["duration"] = {"enabled": True, "trigger": DURATIONS[name]}
        legacy_pool = allowed_pools[0] if len(allowed_pools) == 1 else ("choose" if allowed_pools else "none")
        system.update({
            "tier": minimum_tier, "category": "origin" if catalog == "origin" else "genre-progression",
            "archived": False, "activation": ACTIVATIONS[name], "pool": legacy_pool,
            "cost": {"amount": amount, "scalable": scalable, "ignoresEdge": False, "allowedPools": allowed_pools},
            "roll": roll, "rollModifier": 0, "attackModifier": 0, "damage": damage,
            "woundSeverity": "none", "range": ability_range, "targetMode": target_mode,
            "sourceFocusUuid": "", "sourceNodeId": "",
        })
        metadata = {
            "catalog": catalog, "genres": genres, "progressionBand": band,
            "minimumSuperheroRank": minimum_rank,
            "sourceOccurrences": [
                {"genre": variant["genre"], "band": variant["band"], "page": variant["page"]}
                for variant in variants
            ],
        }
        item = foundry_item(name, item_id, system, sort * 1000, metadata)
        abilities.append(item)
        by_name[name] = item

    catalogs = {}
    for genre, mid, high in (
        ("fantasy", FANTASY_MID, FANTASY_HIGH),
        ("scienceFiction", SCIENCE_FICTION_MID, SCIENCE_FICTION_HIGH),
    ):
        entries = []
        for band, minimum_tier, names in (("mid-tier", 3, mid), ("high-tier", 6, high)):
            for name in names:
                ability = by_name[name]
                entries.append({
                    "id": stable_id("genre-relation", f"{genre}:{band}:{ability['system']['slug']}"),
                    "abilityUuid": f"{ABILITY_PREFIX}{ability['_id']}",
                    "minimumTier": minimum_tier, "catalog": "progression", "minimumSuperheroRank": 0,
                    "snapshot": snapshot(ability),
                })
        catalogs[genre] = {"catalog": "progression", "entries": entries}
    catalogs["superhero"] = {
        "catalog": "origin", "eligibility": "genre", "typeWhitelist": [],
        "entries": [{
            "id": stable_id("genre-relation", f"superhero:origin:{by_name[name]['system']['slug']}"),
            "abilityUuid": f"{ABILITY_PREFIX}{by_name[name]['_id']}",
            "minimumTier": 1, "catalog": "origin",
            "minimumSuperheroRank": 2 if name == "Armored Body" else 0,
            "snapshot": snapshot(by_name[name]),
        } for name in ORIGIN],
    }
    entries = []
    for band, minimum_tier, names in (("mid-tier", 3, REAL_WORLD_MID), ("high-tier", 6, REAL_WORLD_HIGH)):
        for name in names:
            ability = by_name[name]
            entries.append({
                "id": stable_id("genre-relation", f"realWorld:{band}:{ability['system']['slug']}"),
                "abilityUuid": f"{ABILITY_PREFIX}{ability['_id']}",
                "minimumTier": minimum_tier, "catalog": "progression", "minimumSuperheroRank": 0,
                "snapshot": snapshot(ability),
            })
    catalogs["realWorld"] = {"catalog": "progression", "entries": entries}

    return {
        "$schema": "./schema.json", "schemaVersion": SCHEMA_VERSION,
        "source": {
            "kind": "official-reference-document-extraction", "title": BOOK,
            "file": document_path.name, "sha256": hashlib.sha256(document_path.read_bytes()).hexdigest(),
            "license": LICENSE, "note": "Candidate content only; requires in-Foundry review before promotion.",
        },
        "target": {"systemId": "cypherv2", "futurePackId": "genre-abilities"},
        "inventory": {
            "progression": {
                "realWorld": {"midTier": REAL_WORLD_MID, "highTier": REAL_WORLD_HIGH},
                "fantasy": {"midTier": FANTASY_MID, "highTier": FANTASY_HIGH},
                "scienceFiction": {"midTier": SCIENCE_FICTION_MID, "highTier": SCIENCE_FICTION_HIGH},
            },
            "origin": ORIGIN,
        },
        "counts": {"abilities": len(abilities), "progression": len(PROGRESSION_ORDER) + 2, "origin": len(ORIGIN), "relations": sum(len(c["entries"]) for c in catalogs.values())},
        "catalogs": catalogs,
        "review": {
            "status": "candidate",
            "sameNameCollisions": [
                {"name": name, "classification": classification, "existingPacks": packs, "note": note, "reused": False}
                for name, classification, packs, note in COLLISIONS
            ],
            "manualTreatment": sorted(CONDITIONAL_COSTS),
            "originEligibility": "Shared Superhero Genre catalog; no Type whitelist.",
        },
        "abilities": abilities,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--document", type=Path, default=Path("_reference/Cypher-Reference-Document-2026-07-29.docx"))
    parser.add_argument("--output", type=Path, default=Path("content/genre-abilities/candidate-pack.json"))
    args = parser.parse_args()
    candidate = build_candidate(args.document)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(candidate, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(candidate["counts"], indent=2))


if __name__ == "__main__":
    main()

"""Generate beta.3 candidate Type and Type Ability Items from the 2026 CRD DOCX.

This developer-only extractor never writes Compendium packs or system.json.  It
creates the review source at content/types/candidate-pack.json; promotion is a
separate, deliberately manual phase.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import unicodedata
from pathlib import Path

from docx import Document


SCHEMA_VERSION = 1
BOOK = "Cypher Reference Document (2026)"
LICENSE = "2026 Cypher Open License"
ABILITY_PREFIX = "Compendium.cypherv2.type-abilities.Item."
TYPE_ICON = "systems/cypherv2/assets/icons/cyphertype.png"
ABILITY_ICON = "systems/cypherv2/assets/icons/cypherability.png"
TYPE_NAMES = [
    "Barbarian", "Bard", "Cleric", "Druid", "Fighter", "Mage", "Monk",
    "Necromancer", "Paladin", "Ranger", "Rogue", "Archer", "Axe Fighter",
    "Knife Fighter", "Priest", "Sorcerer", "Sword Fighter", "Thief",
    "Two-Weapon Fighter", "Witch", "Burglar", "Noble Warrior",
    "Swashbuckler", "Warrior", "Wizard", "Diplomat", "Engineer", "Medic",
    "Operative", "Pilot", "Soldier", "Android", "Noble", "Psion",
    "Scoundrel", "Starpilot", "Tech", "Trader", "Dealer", "Heavy",
    "Survivor", "Tender", "Crimefighter", "Vigilante", "Enhanced Hero",
    "Powerstar", "Superhuman", "Powerhouse", "Living God",
]
SUPERHERO_DATA = {
    "Crimefighter": (1, 2, 0),
    "Vigilante": (1, 2, 0),
    "Enhanced Hero": (2, 3, 2),
    "Powerstar": (2, 3, 2),
    "Superhuman": (3, 4, 4),
    "Powerhouse": (4, 5, 6),
    "Living God": (5, 6, 8),
}
NUMBER_WORDS = {
    "one": 1, "two": 2, "three": 3, "four": 4, "five": 5,
    "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
}

ASSIGNMENT_NOTES = {
    ("Soldier", "Expert Combatant"):
        "CRD Soldier guidance recommends genre-appropriate weapons, including laser pistols, rail guns, and specific spacecraft weapons.",
    ("Heavy", "Expert Combatant"):
        "CRD Heavy guidance recommends genre-appropriate weapons such as pistols and rifles.",
    ("Tender", "Inspiring Suggestion"):
        "CRD Tender wording omits the Diplomat/Superhuman restriction that the suggested action occur on the ally's next turn and adds an introductory explanatory sentence.",
    ("Crimefighter", "Super Combatant"):
        "CRD Crimefighter wording permits any specific attack (examples include eye lasers, punches, or guns), rather than only a weapon attack.",
    ("Vigilante", "Super Combatant"):
        "CRD Vigilante wording says specific weapon attack, then recommends unarmed attacks for this Type.",
    ("Superhuman", "Super Combatant"):
        "CRD wording limits the chosen training to a specific weapon attack.",
    ("Powerhouse", "Super Combatant"):
        "CRD wording limits the chosen training to a specific weapon attack.",
    ("Living God", "Super Combatant"):
        "CRD wording limits the chosen training to a specific weapon attack.",
    ("Powerhouse", "Enhanced Energy"):
        "CRD Powerhouse occurrence omits the Enabler classification printed in the otherwise identical Powerstar occurrence.",
}


def normalized_space(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def slugify(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def stable_id(kind: str, key: str) -> str:
    return hashlib.sha256(f"cypherv2-beta3:{kind}:{key}".encode("utf-8")).hexdigest()[:16]


def provenance() -> dict:
    return {
        "kind": "other", "sourceUuid": "", "instanceId": "", "grantId": "",
        "status": "active", "contentUuid": "", "contentKey": "",
        "replacement": {
            "active": False, "originalName": "", "originalContentUuid": "",
            "originalContentKey": "", "replacementName": "",
            "replacementContentUuid": "", "replacementContentKey": "",
            "selectionKind": "none",
        },
    }


def package_instance() -> dict:
    return {
        "sourceUuid": "", "instanceId": "", "role": "primary", "attachedAt": 0,
        "selections": {
            "edgePool": "none", "superheroicsPool": "none", "powerShifts": [],
            "poolChoices": [], "skillChoices": [], "abilityChoices": [],
            "suppressedGrantIds": [],
        },
        "parent": provenance(),
    }


def base_system(slug: str, description: str, page: int, tags: list[str]) -> dict:
    return {
        "schemaVersion": 1,
        "slug": slug,
        "description": description,
        "source": {"uuid": "", "book": BOOK, "page": str(page), "license": LICENSE},
        "automation": {
            "mode": "descriptive",
            "duration": {"enabled": False, "trigger": "recovery"},
            "rollDefaults": {},
        },
        "ruleElements": [],
        "tags": tags,
        "grantedBy": provenance(),
    }


def foundry_item(name: str, item_type: str, item_id: str, img: str, system: dict, sort: int) -> dict:
    return {
        "name": name, "type": item_type, "_id": item_id, "img": img,
        "system": system, "effects": [], "folder": None, "sort": sort,
        "ownership": {"default": 0}, "flags": {},
        "_stats": {
            "coreVersion": "14.360", "systemId": None, "systemVersion": None,
            "createdTime": None, "modifiedTime": None, "lastModifiedBy": None,
            "compendiumSource": None, "duplicateSource": None, "exportSource": None,
        },
        "_key": f"!items!{item_id}",
    }


def paragraph_page_numbers(paragraphs) -> list[int]:
    page = 1
    result = []
    for paragraph in paragraphs:
        properties = paragraph._p.pPr
        if properties is not None and properties.pageBreakBefore is not None:
            page += 1
        result.append(page)
        page += len(paragraph._p.xpath(".//w:lastRenderedPageBreak"))
        page += len(paragraph._p.xpath('.//w:br[@w:type="page"]'))
    return result


def runs_html(paragraph, remove_prefix: str = "") -> str:
    remaining = len(remove_prefix)
    fragments = []
    for run in paragraph.runs:
        text = run.text
        if remaining:
            consumed = min(remaining, len(text))
            text = text[consumed:]
            remaining -= consumed
        if not text:
            continue
        value = html.escape(text.replace("\u2028", " ").replace("\u2029", " ")).replace("\n", "<br>")
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        fragments.append(value)
    return "".join(fragments).strip()


def paragraphs_html(paragraphs, first_prefix: str = "") -> str:
    fragments = []
    for index, paragraph in enumerate(paragraphs):
        body = runs_html(paragraph, first_prefix if index == 0 else "")
        if not body:
            continue
        if paragraph.style.name == "Callout":
            fragments.append(f'<aside class="type-ability-callout"><p>{body}</p></aside>')
        else:
            fragments.append(f"<p>{body}</p>")
    return "".join(fragments)


def list_html(paragraphs) -> str:
    items = [f"<li><p>{runs_html(paragraph)}</p></li>" for paragraph in paragraphs if paragraph.text.strip()]
    return f"<ul>{''.join(items)}</ul>" if items else ""


def leading_bold(paragraph) -> str:
    parts = []
    for run in paragraph.runs:
        if run.bold:
            parts.append(run.text)
        elif parts:
            break
    return "".join(parts).strip()


def parse_ability_heading(prefix: str) -> tuple[str, str]:
    heading = normalized_space(prefix[:-1])
    cost_match = re.search(r"\s*\(([^()]*)\)\s*$", heading)
    cost_text = cost_match.group(1) if cost_match else ""
    name = heading[:cost_match.start()].strip() if cost_match else heading
    return name, cost_text


def is_ability_start(paragraph) -> bool:
    if paragraph.style.name not in ("First Paragraph", "Body Text"):
        return False
    prefix = leading_bold(paragraph)
    if not prefix.endswith(":"):
        return False
    lowered = normalized_space(prefix).lower()
    return not lowered.startswith(("at tier", "effort"))


def parse_cost(cost_text: str) -> tuple[int, bool, list[str], str]:
    clean = normalized_space(cost_text)
    match = re.match(r"^(\d+)(\+)?\s+(Might|Speed|Intellect)(?:\s+or\s+(Might|Speed|Intellect))?", clean, re.I)
    if not match:
        return 0, False, [], clean
    pools = [match.group(3).lower()]
    if match.group(4):
        pools.append(match.group(4).lower())
    pools = [pool for pool in ("might", "speed", "intellect") if pool in pools]
    remainder = clean[match.end():].strip()
    return int(match.group(1)), bool(match.group(2)), pools, remainder


def activation_for(paragraphs) -> str:
    normalized = normalized_space(" ".join(paragraph.text for paragraph in paragraphs)).lower()
    if "action to attack; last action" in normalized and "enabler" in normalized:
        return "special"
    if re.search(r"\b(?:ten minutes?|one hour|ten hours?) to (?:build|brew|prepare|complete|craft|initiate)", normalized):
        return "timed"
    if "first action" in normalized:
        return "firstAction"
    if "last action" in normalized:
        return "lastAction"
    if any(re.search(r"\baction\.$", normalized_space(paragraph.text).lower()) for paragraph in paragraphs):
        return "action"
    if re.search(r"\benabler\.?($|\s)", normalized):
        return "enabler"
    if "perpetual" in normalized:
        return "perpetual"
    if re.search(r"\breaction\.?($|\s)", normalized):
        return "reaction"
    return "passive"


def extract_sections(document: Document) -> tuple[list[dict], list[int]]:
    paragraphs = document.paragraphs
    pages = paragraph_page_numbers(paragraphs)
    sections = []
    for index, paragraph in enumerate(paragraphs):
        title = paragraph.text.strip()
        if paragraph.style.name != "Type Headers" or not title.endswith(" Abilities"):
            continue
        type_name = title[:-10]
        if type_name not in TYPE_NAMES:
            continue
        end = index + 1
        while end < len(paragraphs):
            candidate = paragraphs[end]
            if candidate.style.name in ("H2", "Heading 2", "Heading 3"):
                break
            if candidate.style.name == "Type Headers" and candidate.text.strip().endswith("Equipment Bundle"):
                break
            end += 1
        sections.append({"name": type_name, "abilityHeader": index, "end": end})
    if [entry["name"] for entry in sections] != TYPE_NAMES:
        raise ValueError("The CRD Type section inventory no longer matches the reviewed 49-Type order.")
    return sections, pages


def extract_occurrences(document: Document, sections: list[dict]) -> list[dict]:
    paragraphs = document.paragraphs
    occurrences = []
    for section in sections:
        started = False
        starts = []
        for index in range(section["abilityHeader"] + 1, section["end"]):
            paragraph = paragraphs[index]
            if normalized_space(paragraph.text).startswith("At tier 1, you gain the following abilities:"):
                started = True
                continue
            if started and is_ability_start(paragraph):
                starts.append(index)
        for position, start in enumerate(starts):
            end = starts[position + 1] if position + 1 < len(starts) else section["end"]
            prefix = leading_bold(paragraphs[start])
            name, cost_text = parse_ability_heading(prefix)
            body = list(paragraphs[start:end])
            occurrences.append({
                "type": section["name"], "name": name, "costText": cost_text,
                "prefix": prefix, "paragraphs": body, "paragraphIndex": start,
            })
    return occurrences


def genre_for(type_name: str) -> str:
    position = TYPE_NAMES.index(type_name)
    if position < 25:
        return "fantasy"
    if position < 42:
        return "scienceFiction"
    return "superhero"


def type_heading_index(paragraphs, section: dict) -> int:
    for index in range(section["abilityHeader"] - 1, -1, -1):
        paragraph = paragraphs[index]
        if paragraph.style.name == "Heading 3":
            heading = re.sub(r"\s+\(Rank\s+\d+\)\s*$", "", normalized_space(paragraph.text))
            # A Word formatting artifact joins the preceding cross-reference
            # "Soldier" to the visible Diplomat heading in the source XML.
            if heading == section["name"] or heading.endswith(section["name"]):
                return index
    raise ValueError(f"Missing Type heading for {section['name']}.")


def type_source_parts(document: Document, section: dict) -> tuple[list, list, list]:
    paragraphs = document.paragraphs
    heading = type_heading_index(paragraphs, section)
    background_header = next(
        index for index in range(heading + 1, section["abilityHeader"])
        if paragraphs[index].style.name == "Type Headers" and paragraphs[index].text.strip() in ("Background", "Background Options")
    )
    intro = [paragraphs[index] for index in range(heading + 1, background_header) if paragraphs[index].text.strip()]
    background = [
        paragraphs[index] for index in range(background_header + 1, section["abilityHeader"])
        if paragraphs[index].style.name == "Compact" and paragraphs[index].text.strip()
    ]
    equipment = []
    for index in range(section["end"], min(len(paragraphs), section["end"] + 4)):
        paragraph = paragraphs[index]
        if paragraph.style.name == "Type Headers" and paragraph.text.strip().endswith("Equipment Bundle"):
            cursor = index + 1
            while cursor < len(paragraphs) and paragraphs[cursor].style.name not in ("H2", "Heading 2", "Heading 3"):
                if paragraphs[cursor].text.strip():
                    equipment.append(paragraphs[cursor])
                cursor += 1
            break
    return intro, background, equipment


def parse_benefits(document: Document, section: dict) -> list[str]:
    paragraphs = document.paragraphs
    benefits = []
    for index in range(section["abilityHeader"] + 1, section["end"]):
        text = normalized_space(paragraphs[index].text)
        if text.startswith("At tier 1, you gain the following abilities:"):
            break
        if paragraphs[index].style.name == "Compact" and text:
            benefits.append(text)
    return benefits


def wound_bonuses(benefits: list[str]) -> dict:
    result = {"minor": 0, "moderate": 0, "major": 0}
    source = next((entry.lower() for entry in benefits if entry.lower().startswith("able to take")), "")
    for word, amount in NUMBER_WORDS.items():
        for severity in result:
            if re.search(rf"\b{word}\s+more\s+{severity}\s+wound", source):
                result[severity] = amount
    return result


def pool_bonuses(benefits: list[str]) -> dict:
    result = {"might": 0, "speed": 0, "intellect": 0}
    for entry in benefits:
        match = re.fullmatch(r"Add \+(\d+) to (Might|Speed|Intellect) Pool", entry, re.I)
        if match:
            result[match.group(2).lower()] = int(match.group(1))
    return result


def category_flags(benefits: list[str], noun: str) -> dict:
    result = {"light": False, "medium": False, "heavy": False}
    line = next((entry.lower() for entry in benefits if "freely use" in entry.lower() and noun in entry.lower()), "")
    if not line or "cannot freely" in line:
        return result
    if f"all {noun}" in line or (noun == "weapons" and "all light and medium weapons" not in line and "all weapons" in line):
        return {"light": True, "medium": True, "heavy": True}
    for category in result:
        result[category] = category in line
    return result


def weapon_family_flags(benefits: list[str]) -> dict:
    lowered = "\n".join(benefits).lower()
    return {
        "axes": "freely use all axes" in lowered,
        "knives": "freely use all knives" in lowered,
        "swords": "freely use all swords" in lowered,
    }


def build_ability_documents(occurrences: list[dict], pages: list[int]) -> tuple[list[dict], dict[str, dict]]:
    by_name = {}
    for occurrence in occurrences:
        by_name.setdefault(occurrence["name"], []).append(occurrence)
    documents = []
    mapping = {}
    for sort, (name, variants) in enumerate(by_name.items()):
        occurrence = variants[0]
        amount, scalable, pools, residual_cost = parse_cost(occurrence["costText"])
        description = paragraphs_html(occurrence["paragraphs"], occurrence["prefix"])
        activation = activation_for(occurrence["paragraphs"])
        slug = slugify(name)
        item_id = stable_id("ability", slug)
        legacy_pool = pools[0] if len(pools) == 1 else ("choose" if pools else "none")
        tags = ["type-ability"]
        if residual_cost:
            tags.append("manual-additional-cost")
        if name in ("Inspiring Suggestion", "Super Combatant", "Enhanced Energy"):
            tags.append("assignment-note-variant")
        system = base_system(slug, description, pages[occurrence["paragraphIndex"]], tags)
        system.update({
            "tier": 1, "category": "type", "archived": False,
            "activation": activation, "pool": legacy_pool,
            "cost": {"amount": amount, "scalable": scalable, "ignoresEdge": False, "allowedPools": pools},
            "roll": "none", "rollModifier": 0, "attackModifier": 0, "damage": 0,
            "woundSeverity": "none", "range": "", "targetMode": "none",
            "sourceFocusUuid": "", "sourceNodeId": "",
        })
        document = foundry_item(name, "ability", item_id, ABILITY_ICON, system, sort * 1000)
        documents.append(document)
        mapping[name] = document
    return documents, mapping


def build_type_documents(document: Document, sections: list[dict], occurrences: list[dict], abilities: dict[str, dict], pages: list[int]) -> list[dict]:
    paragraphs = document.paragraphs
    assignments_by_type = {name: [] for name in TYPE_NAMES}
    for occurrence in occurrences:
        assignments_by_type[occurrence["type"]].append(occurrence)
    documents = []
    for sort, section in enumerate(sections):
        name = section["name"]
        heading = type_heading_index(paragraphs, section)
        intro, background, equipment = type_source_parts(document, section)
        benefits = parse_benefits(document, section)
        genre = genre_for(name)
        slug = slugify(name)
        item_id = stable_id("type", slug)
        system = base_system(slug, paragraphs_html(intro), pages[heading], ["character-type", genre])
        rank, shift_count, superheroics_bonus = SUPERHERO_DATA.get(name, (0, 0, 0))
        grants = []
        for occurrence in assignments_by_type[name]:
            ability = abilities[occurrence["name"]]
            assignment_id = stable_id("assignment", f"{slug}:{ability['system']['slug']}")
            grants.append({
                "id": assignment_id,
                "abilityUuid": f"{ABILITY_PREFIX}{ability['_id']}",
                "notes": ASSIGNMENT_NOTES.get((name, occurrence["name"]), ""),
                "snapshot": {
                    "name": ability["name"], "img": ability["img"],
                    "system": ability["system"],
                },
                "alternatives": [],
            })
        system.update({
            "poolBonuses": pool_bonuses(benefits),
            "woundBonuses": wound_bonuses(benefits),
            "edgeGrant": {"mode": "choice", "pool": "none", "amount": 1},
            "weaponUse": category_flags(benefits, "weapons"),
            "weaponFamilyUse": weapon_family_flags(benefits),
            "armorUse": category_flags(benefits, "armor"),
            "superhero": {
                "rank": rank, "powerShiftCount": shift_count,
                "superheroics": {"enabled": genre == "superhero", "poolBonus": superheroics_bonus},
            },
            "abilityGrants": grants, "abilityChoiceGroups": [],
            "skillGrants": [], "choiceGroups": [],
            "genre": genre, "customGenreId": "",
            "backgroundOptions": list_html(background),
            "equipmentNotes": paragraphs_html(equipment),
            "equipmentBundleUuid": "", "instance": package_instance(),
        })
        documents.append(foundry_item(name, "characterType", item_id, TYPE_ICON, system, sort * 1000))
    return documents


def build_candidate(document_path: Path) -> dict:
    document = Document(document_path)
    sections, pages = extract_sections(document)
    occurrences = extract_occurrences(document, sections)
    ability_documents, ability_mapping = build_ability_documents(occurrences, pages)
    type_documents = build_type_documents(document, sections, occurrences, ability_mapping, pages)
    return {
        "$schema": "./schema.json",
        "schemaVersion": SCHEMA_VERSION,
        "source": {
            "kind": "official-reference-document-extraction",
            "title": BOOK,
            "file": document_path.name,
            "sha256": hashlib.sha256(document_path.read_bytes()).hexdigest(),
            "license": LICENSE,
            "note": "Candidate content only; requires in-Foundry review before promotion.",
        },
        "target": {"systemId": "cypherv2", "typePackId": "types", "abilityPackId": "type-abilities"},
        "counts": {"types": len(type_documents), "abilities": len(ability_documents), "assignments": len(occurrences)},
        "review": {
            "status": "candidate",
            "sameNameDiscrepancies": [
                {
                    "ability": "Inspiring Suggestion",
                    "resolution": "shared-document-with-assignment-note",
                    "note": "Tender omits the next-turn timing restriction printed for Diplomat and Superhuman.",
                },
                {
                    "ability": "Super Combatant",
                    "resolution": "shared-document-with-assignment-note",
                    "note": "Crimefighter permits a broader specific attack; the other four occurrences say weapon attack, with a further unarmed recommendation for Vigilante.",
                },
                {
                    "ability": "Enhanced Energy",
                    "resolution": "shared-document-with-assignment-note",
                    "note": "Powerstar labels the otherwise identical rule Enabler; Powerhouse omits that label.",
                },
            ],
            "manualTreatment": [
                "Origin Superhero Ability choices are intentionally absent until the future genre-abilities content phase.",
                "Unusual costs and effects remain in rich text when the current runtime cannot model them safely.",
                "Tier 3 and Tier 6 Genre choices remain descriptive Type guidance and are not direct Type Ability assignments.",
            ],
        },
        "types": type_documents,
        "abilities": ability_documents,
    }


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        type=Path,
        default=project_root / "_reference" / "Cypher-Reference-Document-2026-07-29.docx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=project_root / "content" / "types" / "candidate-pack.json",
    )
    args = parser.parse_args()
    candidate = build_candidate(args.source)
    if candidate["counts"] != {"types": 49, "abilities": 106, "assignments": 155}:
        raise ValueError(f"Unexpected candidate counts: {candidate['counts']}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(candidate, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), **candidate["counts"]}, indent=2))


if __name__ == "__main__":
    main()

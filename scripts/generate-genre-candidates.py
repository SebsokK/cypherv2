"""Generate private beta.3 Genre candidates from the official 2026 CRD.

The four chapter-level Genres are represented by the existing Genre Item type.
Subgenres and rules guidance remain rich text; Ability catalogs reuse the
already validated Genre Ability candidate relations.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import unicodedata
from pathlib import Path

from docx import Document


SCHEMA_VERSION = 1
BOOK = "Cypher Reference Document (2026)"
LICENSE = "2026 Cypher Open License"
GENRE_ICON = "systems/cypherv2/assets/icons/cyphergenre.png"

GENRE_INVENTORY = [
    {
        "key": "realWorld",
        "name": "The Real World",
        "heading": "The Real World Genre",
        "pages": "27–38",
        "subgenres": ["Espionage", "Crime Thriller", "Action", "Rescue", "Historical", "Real-World Horror", "Modern Magic"],
        "progressionCatalog": "realWorld",
        "originCatalog": False,
    },
    {
        "key": "fantasy",
        "name": "Fantasy",
        "heading": "Fantasy Genre",
        "pages": "41–105",
        "subgenres": ["Dungeon Fantasy", "Swords & Sorcery", "Epic Fantasy"],
        "progressionCatalog": "fantasy",
        "originCatalog": False,
    },
    {
        "key": "scienceFiction",
        "name": "Science Fiction",
        "heading": "Science Fiction Genre",
        "pages": "105–164",
        "subgenres": ["Hard Science Fiction", "Space Opera", "Postapocalypse"],
        "progressionCatalog": "scienceFiction",
        "originCatalog": False,
    },
    {
        "key": "superhero",
        "name": "Superheroes",
        "heading": "Superheroes Genre",
        "pages": "165–194",
        "subgenres": [],
        "progressionCatalog": "fantasy-and-science-fiction",
        "originCatalog": True,
    },
]


def normalized_space(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def slugify(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def stable_id(kind: str, key: str) -> str:
    return hashlib.sha256(f"cypherv2-beta3:{kind}:{key}".encode("utf-8")).hexdigest()[:16]


def provenance() -> dict:
    return {
        "kind": "other", "sourceUuid": "", "instanceId": "", "grantId": "",
        "status": "active", "contentUuid": "", "contentKey": "",
        "replacement": {
            "active": False, "originalName": "", "originalContentUuid": "",
            "originalContentKey": "", "replacementName": "",
            "replacementContentUuid": "", "replacementContentKey": "",
            "selectionKind": "none",
        },
    }


def paragraph_page_numbers(paragraphs) -> list[int]:
    page = 1
    result = []
    for paragraph in paragraphs:
        properties = paragraph._p.pPr
        if properties is not None and properties.pageBreakBefore is not None:
            page += 1
        result.append(page)
        page += len(paragraph._p.xpath(".//w:lastRenderedPageBreak"))
        page += len(paragraph._p.xpath('.//w:br[@w:type="page"]'))
    return result


def runs_html(paragraph) -> str:
    fragments = []
    for run in paragraph.runs:
        if not run.text:
            continue
        value = html.escape(run.text.replace("\u2028", " ").replace("\u2029", " ")).replace("\n", "<br>")
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        fragments.append(value)
    return "".join(fragments).strip()


def paragraphs_html(paragraphs) -> str:
    fragments: list[str] = []
    list_items: list[str] = []

    def flush_list() -> None:
        if list_items:
            fragments.append(f"<ul>{''.join(list_items)}</ul>")
            list_items.clear()

    for paragraph in paragraphs:
        body = runs_html(paragraph)
        if not body:
            continue
        style = paragraph.style.name
        if style == "Compact":
            list_items.append(f"<li>{body}</li>")
            continue
        flush_list()
        if style in ("Heading 1", "Heading 2", "H2"):
            fragments.append(f"<h2>{body}</h2>")
        elif style in ("Heading 3", "Type Headers"):
            fragments.append(f"<h3>{body}</h3>")
        elif style == "Heading 4":
            fragments.append(f"<h4>{body}</h4>")
        elif style == "Callout":
            fragments.append(f'<aside class="genre-callout"><p>{body}</p></aside>')
        else:
            fragments.append(f"<p>{body}</p>")
    flush_list()
    return "".join(fragments)


def find_paragraph(paragraphs, title: str, *, style: str | None = None, after: int = -1) -> int:
    matches = [
        index for index, paragraph in enumerate(paragraphs)
        if index > after and normalized_space(paragraph.text) == title and (style is None or paragraph.style.name == style)
    ]
    if not matches:
        raise ValueError(f"Missing CRD paragraph '{title}' ({style or 'any style'}) after {after}.")
    return matches[0]


def selected_description(document: Document, key: str) -> str:
    paragraphs = document.paragraphs

    def index(title: str, style: str | None = None, after: int = -1) -> int:
        return find_paragraph(paragraphs, title, style=style, after=after)

    def segment(start: int, end: int, include_start: bool = True) -> list:
        return list(paragraphs[start if include_start else start + 1:end])

    if key == "realWorld":
        start = index("The Real World Genre", "Heading 1")
        end = index("Genre Character Abilities from Types and Foci", "Heading 1", start)
        chosen = segment(start, end, include_start=False)
    elif key == "fantasy":
        root = index("Fantasy Genre", "Heading 1")
        dungeon = index("Dungeon Fantasy", "Heading 2", root)
        dungeon_types = index("Dungeon Fantasy Types", "Heading 3", dungeon)
        swords = index("Swords & Sorcery", "Heading 2", dungeon_types)
        swords_types = index("Swords & Sorcery Types", "Heading 3", swords)
        epic = index("Epic Fantasy", "Heading 2", swords_types)
        epic_types = index("Epic Fantasy Types", None, epic)
        genre_abilities = index("Fantasy Genre Abilities", "Heading 2", epic_types)
        mid = index("Mid-Tier Fantasy Abilities", "Heading 3", genre_abilities)
        species = index("Fantasy Species", "Heading 3", mid)
        dragonfolk = index("Dragonfolk", "Type Headers", species)
        chosen = [
            *segment(root, dungeon, include_start=False),
            *segment(dungeon, dungeon_types),
            *segment(swords, swords_types),
            *segment(epic, epic_types),
            *segment(genre_abilities, mid),
            *segment(species, dragonfolk),
        ]
    elif key == "scienceFiction":
        root = index("Science Fiction Genre", "Heading 1")
        hard = index("Hard Science Fiction", "Heading 2", root)
        hard_types = index("Hard Science Fiction Types", "Heading 3", hard)
        space = index("Space Opera", "Heading 2", hard_types)
        space_types = index("Space Opera Types", "Heading 3", space)
        post = index("Postapocalypse", "Heading 2", space_types)
        post_types = index("Postapocalypse Types", "H2", post)
        genre_abilities = index("Science Fiction Genre Abilities", "Heading 2", post_types)
        mid = index("Mid-Tier Science Fiction Abilities", "Heading 3", genre_abilities)
        species = index("Science Fiction Species", "Heading 2", mid)
        first_species = index("Aarak", "Heading 3", species)
        species_or_type = index("Species or Type?", "Heading 3", first_species)
        adapting = index("Adapting Types From Other Genres", "H2", species_or_type)
        superhero = index("Superheroes Genre", "Heading 1", adapting)
        chosen = [
            *segment(root, hard, include_start=False),
            *segment(hard, hard_types),
            *segment(space, space_types),
            *segment(post, post_types),
            *segment(genre_abilities, mid),
            *segment(species, first_species),
            *segment(species_or_type, adapting),
            *segment(adapting, superhero),
        ]
    elif key == "superhero":
        root = index("Superheroes Genre", "Heading 1")
        types = index("Superhero Types", "Heading 3", root)
        impossible = index("Really Impossible Tasks", "H2", types)
        origin = index("Origin Superhero Abilities", "Heading 2", impossible)
        adapting = index("Adapting Alternate Types", "Heading 3", origin)
        foci = index("Foci", "Heading 1", adapting)
        chosen = [
            *segment(root, types, include_start=False),
            *segment(impossible, origin),
            *segment(adapting, foci),
        ]
    else:
        raise ValueError(f"Unsupported Genre key: {key}")
    description = paragraphs_html(chosen)
    if not description:
        raise ValueError(f"Generated empty description for {key}.")
    return description


def superhero_progression_catalog(catalogs: dict) -> list[dict]:
    entries = []
    seen_abilities = set()
    for catalog_name in ("fantasy", "scienceFiction"):
        for entry in catalogs[catalog_name]["entries"]:
            if entry["abilityUuid"] in seen_abilities:
                continue
            seen_abilities.add(entry["abilityUuid"])
            entries.append(entry)
    if len(entries) != 41:
        raise ValueError(f"Expected 41 unique Superhero progression options; found {len(entries)}.")
    return entries


def genre_catalog(key: str, ability_source: dict) -> list[dict]:
    catalogs = ability_source["catalogs"]
    if key == "realWorld":
        return json.loads(json.dumps(catalogs["realWorld"]["entries"]))
    if key == "fantasy":
        return json.loads(json.dumps(catalogs["fantasy"]["entries"]))
    if key == "scienceFiction":
        return json.loads(json.dumps(catalogs["scienceFiction"]["entries"]))
    if key == "superhero":
        return json.loads(json.dumps([
            *superhero_progression_catalog(catalogs),
            *catalogs["superhero"]["entries"],
        ]))
    raise ValueError(f"Unsupported Genre key: {key}")


def build_candidate(document_path: Path, ability_candidate_path: Path) -> dict:
    document = Document(document_path)
    ability_source = json.loads(ability_candidate_path.read_text(encoding="utf-8"))
    if ability_source.get("counts") != {"abilities": 69, "progression": 43, "origin": 26, "relations": 74}:
        raise ValueError("Genre Ability candidate counts do not match the validated 43 + 26 architecture.")

    paragraphs = document.paragraphs
    pages = paragraph_page_numbers(paragraphs)
    heading_pages = {}
    for genre in GENRE_INVENTORY:
        heading_index = find_paragraph(paragraphs, genre["heading"], style="Heading 1")
        heading_pages[genre["key"]] = pages[heading_index]

    genres = []
    for sort, inventory in enumerate(GENRE_INVENTORY):
        key = inventory["key"]
        slug = slugify(inventory["name"])
        item_id = stable_id("genre", slug)
        catalog = genre_catalog(key, ability_source)
        system = {
            "schemaVersion": 1,
            "slug": slug,
            "description": selected_description(document, key),
            "source": {"uuid": "", "book": BOOK, "page": inventory["pages"], "license": LICENSE},
            "automation": {"mode": "descriptive", "duration": {"enabled": False, "trigger": "recovery"}, "rollDefaults": {}},
            "ruleElements": [],
            "tags": ["genre", f"genre:{slug}"],
            "grantedBy": provenance(),
            "abilityCatalog": catalog,
            "options": {"totalEffortCapMode": "unlimited" if key == "superhero" else "core"},
            "legacyKey": key,
        }
        genres.append({
            "name": inventory["name"], "type": "genre", "_id": item_id,
            "img": GENRE_ICON, "system": system, "effects": [], "folder": None,
            "sort": sort * 1000, "ownership": {"default": 0}, "flags": {},
            "_stats": {
                "coreVersion": "14.360", "systemId": None, "systemVersion": None,
                "createdTime": None, "modifiedTime": None, "lastModifiedBy": None,
                "compendiumSource": None, "duplicateSource": None, "exportSource": None,
            },
            "_key": f"!items!{item_id}",
        })

    return {
        "$schema": "./schema.json",
        "schemaVersion": SCHEMA_VERSION,
        "source": {
            "kind": "official-reference-document-extraction", "title": BOOK,
            "file": document_path.name, "sha256": hashlib.sha256(document_path.read_bytes()).hexdigest(),
            "license": LICENSE, "note": "Candidate content only; requires in-Foundry review before promotion.",
        },
        "target": {"systemId": "cypherv2", "futurePackId": "genres", "abilityPackId": "genre-abilities"},
        "inventory": GENRE_INVENTORY,
        "counts": {"genres": 4, "progressionRelations": 89, "originRelations": 26, "catalogRelations": 115},
        "review": {
            "status": "candidate",
            "chapterHeadingPages": heading_pages,
            "classification": {
                "structured": ["Genre Ability UUID relations", "progression/origin catalog kind", "minimum Tier", "minimum Superhero Rank", "total Effort cap mode"],
                "descriptive": ["subgenre guidance", "character creation guidance", "skills", "Descriptor/Type/Focus/Species guidance", "equipment and cyphers", "wound treatment", "currency", "Power Shifts and optional rules"],
            },
            "relationshipPolicy": "The Real World references its two Skill-choice progression Abilities. Superheroes reuses the deduplicated union of the Fantasy and Science Fiction progression relations plus all 26 Origin relations. No Ability document is cloned.",
            "manualGuidance": [
                "Subgenre-specific creation, equipment, treatment, and optional rules remain guide-don't-enforce rich text.",
                "Origin choice timing and alternate-Type adaptation remain descriptive; catalog: origin prevents those entries from entering ordinary Genre progression.",
            ],
        },
        "genres": genres,
    }


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=project_root / "_reference" / "Cypher-Reference-Document-2026-07-29.docx")
    parser.add_argument("--abilities", type=Path, default=project_root / "content" / "genre-abilities" / "candidate-pack.json")
    parser.add_argument("--output", type=Path, default=project_root / "content" / "genres" / "candidate-pack.json")
    args = parser.parse_args()
    candidate = build_candidate(args.source, args.abilities)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(candidate, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), **candidate["counts"]}, indent=2))


if __name__ == "__main__":
    main()
